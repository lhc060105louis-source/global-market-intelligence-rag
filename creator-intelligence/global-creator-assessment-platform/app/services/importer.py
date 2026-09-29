import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import COMMERCIAL_WEIGHTS, RISK_WEIGHTS
from app.models import (
    AssessmentInput,
    ImportError,
    ImportJob,
    Kol,
    KolScoreSummary,
    ScoreRecord,
)
from app.assessment import calculate_assessment
from app.identity import find_existing_kol
from app.market import normalize_market
from app.services.scoring import risk_level, summarize_dimensions

HEADER_ALIASES = {
    "platform": "platform",
    "Country": "country",
    "handle": "handle",
    "platform account id": "platform_account_id",
    "Name": "name",
    "profile": "profile_url",
    "language": "language",
    "content categories": "content_categories",
    "followers": "followers",
    "engagement rate": "average_engagement_rate",
    "audience country share": "audience_country_ratio",
    "audience fit": "audience_fit",
    "content relevance and expertise": "content_relevance",
    "interaction quality": "interaction_quality",
    "voice of customer value": "voc_value",
    "commercial efficiency": "commercial_efficiency",
    "brand fit": "brand_fit",
    "partnership executability": "execution_capability",
    "historical controversy": "historical_controversy",
    "advertising disclosure compliance risk": "ad_disclosure",
    "competitor conflict": "competitor_conflict",
    "fake traffic risk": "fake_traffic",
    "data and privacy risk": "data_privacy",
    "minors and sensitive audience risk": "sensitive_audience",
    "sustainability and technical claims risk": "sustainability_claims",
    "partnership execution risk": "execution_risk",
}

COMPLETE_COMMERCIAL_HEADERS = {
    "KOLName": "name",
    "influencerName": "name",
    "Primary Market (DE/GB/FR/MULTI)": "country",
    "Content Direction (review/ev/luxury/family/tech/lifestyle)": "contentDirection",
    "Target Brand Type (premium/mainstream/ev-brand)": "targetBrand",
    "Target Market Audience Share %": "geo", "Language Proficiency (1/2/3)": "lang",
    "Automotive Interest Audience %": "autoInterest", "Audience Aged 25-55 %": "age",
    "Income Level (high/mid/low)": "income", "Automotive Content Focus %": "focus",
    "Review Depth (deep/mid/surface)": "depth", "Professional Credibility (high/mid/low)": "credibility",
    "ERR % (YouTube API Auto-collected)": "err", "Video Completion Rate %": "completion",
    "Comment Quality (high/mid/low)": "commentQuality", "Share-to-Save Ratio %": "shareSave",
    "VOC Topic Depth (high/mid/low)": "vocDepth", "VOC Negative Sentiment Detection (high/mid/low)": "vocNeg",
    "Historical Owner Feedback (yes/sometimes/no)": "vocHistory", "Quoted CPM €": "cpm",
    "Industry Benchmark CPM €": "benchCpm", "Content Reuse Rights (full/limited/none)": "reuse",
    "Exclusivity Requirement (none/soft/hard)": "exclusive", "Brand Tone Alignment (match/neutral/conflict)": "brandTone",
    "Historical Partnership Tone (match/neutral/conflict)": "histTone", "Content Style Consistency (high/mid/low)": "styleConsist",
    "Fulfillment Rate %": "fulfill", "Brief Cooperation (high/mid/low)": "briefCoop",
    "Data Review Willingness (active/passive/refuse)": "dataReady",
}

COMPLETE_RISK_HEADERS = {
    "KOLName": "name", "Primary Market (DE/GB/FR/MULTI)": "country",
    "Major Negative Event (none/minor/serious/critical)": "incident", "False Advertising Record (none/minor/serious)": "falsead",
    "Public Sentiment Reach (none/local/wide)": "sentiment", "Advertising Label Practice (always/sometimes/never)": "adlabel",
    "Platform Regulatory Action (none/warning/penalty)": "penalty", "Compliance Willingness (high/mid/low)": "compliance",
    "Competitor Affiliation Status (none/nonexclusive/exclusive/ambassador)": "competitor", "Recent Competitor Content %": "compcontentpct",
    "Competitor Brand Tier (none/indirect/direct)": "complevel", "Fake Follower Share %": "fakepct",
    "Follower Spike History (none/once/multiple)": "spikegrowth", "Template-like Comments (normal/some/heavy)": "templatecomment",
    "GDPR Violation Record (none/minor/serious)": "gdpr", "Data Use Practice (compliant/unclear/violation)": "datause",
    "Minor Audience Share %": "minorpct", "Age Suitability (suitable/partial/unsuitable)": "agesuit",
    "Historical Exaggerated Claims (none/minor/serious)": "exaggerate", "Autonomous Driving Claims Risk (none/cautious/exaggerated)": "adas",
    "Technical Accuracy (high/mid/low)": "techaccuracy", "Historical Late Deletion (none/occasional/frequent)": "latedelete",
    "Brief Revision Cooperation (cooperative/friction/refuse)": "briefreject",
}

BASE_FIELDS = {
    "name",
    "platform",
    "platform_account_id",
    "handle",
    "profile_url",
    "country",
    "language",
    "content_categories",
    "followers",
    "average_engagement_rate",
    "audience_country_ratio",
}


def normalize_header(value: object) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    return HEADER_ALIASES.get(lowered, HEADER_ALIASES.get(text, lowered))


def _rows_from_csv(content: bytes) -> list[dict[str, Any]]:
    reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
    if not reader.fieldnames:
        raise ValueError("file has no header row")
    return [
        {normalize_header(key): value for key, value in row.items() if key is not None}
        for row in reader
    ]


def _rows_from_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        headers = [normalize_header(value) for value in next(values)]
    except StopIteration as exc:
        raise ValueError("file has no header row") from exc
    return [dict(zip(headers, row, strict=False)) for row in values]


def _normalized_sheet_name(value: str) -> str:
    return "".join(value.split()).lower()


def _complete_rows(content: bytes) -> list[dict[str, Any]] | None:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets = {_normalized_sheet_name(sheet.title): sheet for sheet in workbook.worksheets}
    commercial = sheets.get("commercialvaluemodel")
    risk = sheets.get("riskassessmentmodel")
    if commercial is None or risk is None:
        return None

    def mapped_rows(sheet, mapping):
        values = sheet.iter_rows(values_only=True)
        headers = next(values, ())
        keys = [mapping.get(str(header or "").strip()) for header in headers]
        return [
            {key: value for key, value in zip(keys, row, strict=False) if key}
            for row in values if any(value is not None for value in row)
        ]

    commercial_rows = mapped_rows(commercial, COMPLETE_COMMERCIAL_HEADERS)
    risk_by_name = {
        str(row.get("name") or "").strip().casefold(): row
        for row in mapped_rows(risk, COMPLETE_RISK_HEADERS) if row.get("name")
    }
    rows = []
    for row in commercial_rows:
        name = _clean_text(row.get("name"))
        joined = dict(row)
        joined["_risk_inputs"] = risk_by_name.get((name or "").casefold(), {})
        joined["_complete_package"] = True
        rows.append(joined)
    return rows


def read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return _rows_from_csv(content)
    if extension == ".xlsx":
        return _complete_rows(content) or _rows_from_xlsx(content)
    raise ValueError("unsupported file type")


def _clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _clean_number(value: object, *, integer: bool = False) -> float | int | None:
    text = _clean_text(value)
    if text is None:
        return None
    number = float(text.replace(",", ""))
    return int(number) if integer else number


def _clean_ratio(value: object) -> float | None:
    text = _clean_text(value)
    if text is None:
        return None
    if text.endswith("%"):
        return float(text[:-1].replace(",", "")) / 100
    number = float(text.replace(",", ""))
    return number / 100 if number > 1 else number


def _clean_score(value: object) -> float | None:
    score = _clean_number(value)
    if score is None:
        return None
    if not 0 <= score <= 100:
        raise ValueError("score must be between 0 and 100")
    return float(score)


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    if row.get("_complete_package"):
        name = _clean_text(row.get("name"))
        country = _clean_text(row.get("country"))
        if not name:
            raise ValueError("KOL name is required")
        if not country:
            raise ValueError("country is required")
        if country.upper() not in {"DE", "GB", "FR", "MULTI"}:
            raise ValueError("country must be DE, GB, FR, or MULTI")
        commercial = {
            key: (None if str(value or "").strip() == "← YouTube API" else value)
            for key, value in row.items()
            if key not in {"name", "country", "_risk_inputs", "_complete_package"}
        }
        commercial["contractFlex"] = None
        risk = {
            key: value for key, value in row["_risk_inputs"].items()
            if key not in {"name", "country"}
        }
        return {
            "platform": "YouTube", "country": country.upper(), "handle": name,
            "platform_account_id": None, "name": name, "profile_url": None,
            "language": None, "content_categories": _clean_text(row.get("contentDirection")),
            "followers": None, "average_engagement_rate": None,
            "audience_country_ratio": _clean_ratio(row.get("geo")),
            "_commercial_inputs": commercial, "_risk_inputs": risk,
        }
    platform = _clean_text(row.get("platform"))
    country_text = _clean_text(row.get("country"))
    handle = _clean_text(row.get("handle"))
    account_id = _clean_text(row.get("platform_account_id"))
    if not platform:
        raise ValueError("platform is required")
    if not country_text:
        raise ValueError("country is required")
    country = normalize_market(country_text)
    if not handle and not account_id:
        raise ValueError("handle or platform_account_id is required")

    cleaned = {
        "platform": platform,
        "country": country,
        "handle": handle,
        "platform_account_id": account_id,
        "name": _clean_text(row.get("name")),
        "profile_url": _clean_text(row.get("profile_url")),
        "language": _clean_text(row.get("language")),
        "content_categories": _clean_text(row.get("content_categories")),
        "followers": _clean_number(row.get("followers"), integer=True),
        "average_engagement_rate": _clean_ratio(row.get("average_engagement_rate")),
        "audience_country_ratio": _clean_ratio(row.get("audience_country_ratio")),
    }
    for dimension in COMMERCIAL_WEIGHTS | RISK_WEIGHTS:
        cleaned[dimension] = _clean_score(row.get(dimension))
    return cleaned


def _find_existing(session: Session, row: dict[str, Any]) -> Kol | None:
    return find_existing_kol(
        session, platform=row["platform"], platform_account_id=row.get("platform_account_id"),
        profile_url=row.get("profile_url"), handle=row.get("handle"),
    )


def _upsert_scores(
    session: Session,
    kol: Kol,
    row: dict[str, Any],
    source: str,
) -> None:
    for score_type, weights in (
        ("commercial", COMMERCIAL_WEIGHTS),
        ("risk", RISK_WEIGHTS),
    ):
        for dimension in weights:
            score = row.get(dimension)
            if score is None:
                continue
            record = session.scalar(
                select(ScoreRecord).where(
                    ScoreRecord.kol_id == kol.id,
                    ScoreRecord.score_type == score_type,
                    ScoreRecord.dimension == dimension,
                )
            )
            if record is None:
                record = ScoreRecord(
                    kol_id=kol.id,
                    score_type=score_type,
                    dimension=dimension,
                )
                session.add(record)
            record.auto_score = score
            record.source = source


def refresh_summary(session: Session, kol: Kol) -> KolScoreSummary:
    records = list(
        session.scalars(select(ScoreRecord).where(ScoreRecord.kol_id == kol.id))
    )
    grouped: dict[str, dict[str, object]] = {"commercial": {}, "risk": {}}
    for record in records:
        grouped[record.score_type][record.dimension] = {
            "auto": record.auto_score,
            "manual": record.manual_score,
        }

    commercial = summarize_dimensions(grouped["commercial"], COMMERCIAL_WEIGHTS)
    risk = summarize_dimensions(grouped["risk"], RISK_WEIGHTS)
    summary = session.scalar(
        select(KolScoreSummary).where(KolScoreSummary.kol_id == kol.id)
    )
    if summary is None:
        summary = KolScoreSummary(kol_id=kol.id)
        session.add(summary)
    summary.commercial_score = commercial.score
    summary.commercial_completeness = commercial.completeness
    summary.commercial_status = commercial.status
    summary.risk_score = risk.score
    summary.risk_completeness = risk.completeness
    summary.risk_status = risk.status
    summary.risk_level = risk_level(risk.score)
    return summary


def persist_assessment(
    session: Session,
    kol: Kol,
    commercial_inputs: dict[str, Any],
    risk_inputs: dict[str, Any],
    *,
    source: str | None,
) -> AssessmentInput:
    """Persist raw inputs and refresh every calculable automatic dimension."""
    result = calculate_assessment(commercial_inputs, risk_inputs)
    raw = session.get(AssessmentInput, kol.id)
    if raw is None:
        raw = AssessmentInput(kol_id=kol.id)
        session.add(raw)
    raw.commercial_inputs = dict(commercial_inputs)
    raw.risk_inputs = dict(risk_inputs)
    raw.flags = list(result.flags)
    if source is not None:
        raw.source = source

    for score_type, dimensions in (
        ("commercial", result.commercial_dimensions),
        ("risk", result.risk_dimensions),
    ):
        for dimension, calculated in dimensions.items():
            record = session.scalar(
                select(ScoreRecord).where(
                    ScoreRecord.kol_id == kol.id,
                    ScoreRecord.score_type == score_type,
                    ScoreRecord.dimension == dimension,
                )
            )
            if record is None:
                record = ScoreRecord(
                    kol_id=kol.id,
                    score_type=score_type,
                    dimension=dimension,
                )
                session.add(record)
            record.auto_score = calculated.score
            if source is not None:
                record.source = source
                record.evidence = "\n".join(calculated.evidence) or None
    session.flush()
    # Summary is derived from persisted records so manual scores override automatic
    # scores and unavailable dimensions remain excluded from the denominator.
    refresh_summary(session, kol)
    return raw


def import_file(
    session: Session,
    filename: str,
    content: bytes,
) -> ImportJob:
    rows = read_rows(filename, content)
    safe_filename = Path(filename.replace("\\", "/")).name
    job = ImportJob(filename=safe_filename, status="processing", total_rows=len(rows))
    session.add(job)
    session.flush()

    for row_number, raw_row in enumerate(rows, start=2):
        try:
            created = False
            with session.begin_nested():
                row = _normalize_row(raw_row)
                kol = _find_existing(session, row)
                if kol is None:
                    kol = Kol(**{key: row[key] for key in BASE_FIELDS})
                    session.add(kol)
                    session.flush()
                    created = True
                else:
                    for key in BASE_FIELDS:
                        if row[key] is not None:
                            setattr(kol, key, row[key])
                _upsert_scores(session, kol, row, f"import:{safe_filename}")
                if "_commercial_inputs" in row:
                    persist_assessment(
                        session, kol, row["_commercial_inputs"], row["_risk_inputs"],
                        source=f"import:{safe_filename}",
                    )
                session.flush()
                refresh_summary(session, kol)
            if created:
                job.created_count += 1
            else:
                job.updated_count += 1
        except (TypeError, ValueError) as exc:
            job.failed_count += 1
            job.errors.append(
                ImportError(row_number=row_number, message=str(exc))
            )

    job.status = "completed"
    session.commit()
    session.refresh(job)
    return job
