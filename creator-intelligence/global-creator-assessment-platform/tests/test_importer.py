from io import BytesIO

from openpyxl import Workbook


COMPLETE_COMMERCIAL_HEADERS = [
    "KOLName", "Primary Market (DE/GB/FR/MULTI)",
    "Content Direction (review/ev/luxury/family/tech/lifestyle)",
    "Target Brand Type (premium/mainstream/ev-brand)", "Target Market Audience Share %",
    "Language Proficiency (1/2/3)", "Automotive Interest Audience %", "Audience Aged 25-55 %",
    "Income Level (high/mid/low)", "Automotive Content Focus %", "Review Depth (deep/mid/surface)",
    "Professional Credibility (high/mid/low)", "ERR % (YouTube API Auto-collected)", "Video Completion Rate %",
    "Comment Quality (high/mid/low)", "Share-to-Save Ratio %", "VOC Topic Depth (high/mid/low)",
    "VOC Negative Sentiment Detection (high/mid/low)", "Historical Owner Feedback (yes/sometimes/no)", "Quoted CPM €",
    "Industry Benchmark CPM €", "Content Reuse Rights (full/limited/none)", "Exclusivity Requirement (none/soft/hard)",
    "Brand Tone Alignment (match/neutral/conflict)", "Historical Partnership Tone (match/neutral/conflict)",
    "Content Style Consistency (high/mid/low)", "Fulfillment Rate %", "Brief Cooperation (high/mid/low)",
    "Data Review Willingness (active/passive/refuse)",
]

COMPLETE_RISK_HEADERS = [
    "KOLName", "Primary Market (DE/GB/FR/MULTI)",
    "Major Negative Event (none/minor/serious/critical)", "False Advertising Record (none/minor/serious)",
    "Public Sentiment Reach (none/local/wide)", "Advertising Label Practice (always/sometimes/never)",
    "Platform Regulatory Action (none/warning/penalty)", "Compliance Willingness (high/mid/low)",
    "Competitor Affiliation Status (none/nonexclusive/exclusive/ambassador)", "Recent Competitor Content %",
    "Competitor Brand Tier (none/indirect/direct)", "Fake Follower Share %",
    "Follower Spike History (none/once/multiple)", "Template-like Comments (normal/some/heavy)",
    "GDPR Violation Record (none/minor/serious)", "Data Use Practice (compliant/unclear/violation)",
    "Minor Audience Share %", "Age Suitability (suitable/partial/unsuitable)",
    "Historical Exaggerated Claims (none/minor/serious)", "Autonomous Driving Claims Risk (none/cautious/exaggerated)",
    "Technical Accuracy (high/mid/low)", "Historical Late Deletion (none/occasional/frequent)",
    "Brief Revision Cooperation (cooperative/friction/refuse)",
]


def complete_package_workbook_bytes() -> bytes:
    workbook = Workbook()
    commercial = workbook.active
    commercial.title = "Commercial Value Model"
    commercial.append(COMPLETE_COMMERCIAL_HEADERS)
    commercial_values = {
        "KOLName": "Max Torques", "Primary Market (DE/GB/FR/MULTI)": "GB",
        "Content Direction (review/ev/luxury/family/tech/lifestyle)": "ev",
        "Target Market Audience Share %": 78, "Language Proficiency (1/2/3)": 3,
        "Automotive Interest Audience %": 85, "Audience Aged 25-55 %": 72,
        "Income Level (high/mid/low)": "high", "ERR % (YouTube API Auto-collected)": "← YouTube API",
    }
    commercial.append([commercial_values.get(header) for header in COMPLETE_COMMERCIAL_HEADERS])

    risk = workbook.create_sheet("Risk Assessment Model")
    risk.append(COMPLETE_RISK_HEADERS)
    risk_values = {
        "KOLName": "Max Torques", "Primary Market (DE/GB/FR/MULTI)": "GB",
        "Major Negative Event (none/minor/serious/critical)": "none",
        "Historical Late Deletion (none/occasional/frequent)": "none",
        "Brief Revision Cooperation (cooperative/friction/refuse)": "cooperative",
    }
    risk.append([risk_values.get(header) for header in COMPLETE_RISK_HEADERS])

    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def upload_csv(client, text: str, filename: str = "kols.csv"):
    return client.post(
        "/imports",
        files={"file": (filename, text.encode("utf-8"), "text/csv")},
    )


def test_import_csv_creates_kol_and_scores(client):
    csv_data = (
        "platform,country,handle,name,followers,audience_fit,"
        "historical_controversy\n"
        "YouTube,UK,@alex,Alex EV,120000,85,10\n"
    )

    response = upload_csv(client, csv_data)

    assert response.status_code == 201
    assert response.json()["created"] == 1
    detail = client.get("/kols/1").json()
    assert detail["handle"] == "@alex"
    assert detail["score_summary"]["commercial_score"] == 85
    assert detail["score_summary"]["commercial_completeness"] == 0.2
    assert detail["score_summary"]["risk_score"] == 10


def test_import_updates_duplicate_handle(client):
    first = "platform,country,handle,followers\nYouTube,UK,@alex,100\n"
    second = "platform,country,handle,followers\nYouTube,UK,@alex,200\n"
    upload_csv(client, first, "first.csv")

    response = upload_csv(client, second, "second.csv")

    assert response.json()["updated"] == 1
    kols = client.get("/kols").json()
    assert len(kols) == 1
    assert kols[0]["followers"] == 200


def test_bad_row_does_not_block_valid_row(client):
    csv_data = (
        "platform,country,handle\n"
        "YouTube,UK,@valid\n"
        "TikTok,,@invalid\n"
    )

    response = upload_csv(client, csv_data)

    body = response.json()
    assert body["created"] == 1
    assert body["failed"] == 1
    job = client.get(f"/imports/{body['job_id']}")
    assert job.status_code == 200
    assert job.json()["errors"][0]["row_number"] == 3
    assert "country is required" in job.json()["errors"][0]["message"]


def test_english_headers_and_germany_are_normalized(client):
    csv_data = "platform,Country,handle,followers\nYouTube,Germany,@auto,10000\n"

    response = upload_csv(client, csv_data)

    assert response.status_code == 201
    kol = client.get("/kols").json()[0]
    assert kol["country"] == "DE"


def test_import_xlsx(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["platform", "country", "handle", "followers"])
    sheet.append(["TikTok", "United Kingdom", "@evuk", 5000])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/imports",
        files={
            "file": (
                "kols.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    assert response.json()["created"] == 1


def test_unsupported_extension_returns_415(client):
    response = client.post(
        "/imports",
        files={"file": ("kols.txt", b"x", "text/plain")},
    )

    assert response.status_code == 415


def test_imports_complete_package_workbook_and_maps_raw_assessment_inputs(client):
    response = client.post(
        "/imports",
        files={"file": ("BYD_Xpeng_Creator_Assessment_Data.xlsx", complete_package_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 201
    assert response.json() == {
        "job_id": 1, "total_rows": 1, "created": 1, "updated": 0, "failed": 0
    }
    match = client.get("/kols", params={"keyword": "Max Torques"}).json()[0]
    detail = client.get(f"/kols/{match['id']}").json()
    assert detail["country"] == "GB"
    assert detail["content_categories"] == "ev"
    assert detail["average_engagement_rate"] is None
    assert detail["commercial_inputs"]["err"] is None
    assert detail["commercial_inputs"]["geo"] == 78
    assert detail["commercial_inputs"]["contractFlex"] is None
    assert len(detail["commercial_inputs"]) == 28
    assert detail["risk_inputs"]["incident"] == "none"
    assert detail["risk_inputs"]["briefreject"] == "cooperative"
    assert len(detail["risk_inputs"]) == 21
    assert detail["score_summary"]["commercial_score"] is not None
    assert detail["score_summary"]["risk_score"] is not None


def test_complete_package_bad_row_does_not_block_valid_rows(client):
    workbook = Workbook()
    commercial = workbook.active
    commercial.title = " Commercial Value Model "
    commercial.append(["influencerName", "Primary Market (DE/GB/FR/MULTI)", "Target Market Audience Share %"])
    commercial.append(["Valid", "GB", 70])
    commercial.append(["Broken", None, 50])
    risk = workbook.create_sheet("risk assessment model")
    risk.append(["KOLName", "Primary Market (DE/GB/FR/MULTI)", "Major Negative Event (none/minor/serious/critical)"])
    risk.append(["Valid", "GB", "none"])
    content = BytesIO()
    workbook.save(content)

    response = client.post("/imports", files={"file": ("complete.xlsx", content.getvalue())})

    assert response.status_code == 201
    assert response.json()["created"] == 1
    assert response.json()["failed"] == 1


def test_assessment_failure_rolls_back_whole_row_and_next_row_succeeds(client):
    workbook = Workbook()
    commercial = workbook.active
    commercial.title = "Commercial Value Model"
    commercial.append([
        "KOLName", "Primary Market (DE/GB/FR/MULTI)", "Quoted CPM €", "Industry Benchmark CPM €",
        "Content Reuse Rights (full/limited/none)", "Exclusivity Requirement (none/soft/hard)",
    ])
    commercial.append(["Atomic Failure", "GB", 10, 0, "full", "none"])
    commercial.append(["Still Valid", "GB", None, None, None, None])
    risk = workbook.create_sheet("Risk Assessment Model")
    risk.append(["KOLName", "Primary Market (DE/GB/FR/MULTI)", "Fake Follower Share %"])
    risk.append(["Atomic Failure", "GB", "not-a-number"])
    risk.append(["Still Valid", "GB", None])
    content = BytesIO()
    workbook.save(content)

    response = client.post("/imports", files={"file": ("atomic.xlsx", content.getvalue())})

    assert response.status_code == 201
    assert response.json()["created"] == 1
    assert response.json()["failed"] == 1
    assert [kol["name"] for kol in client.get("/kols").json()] == ["Still Valid"]


def test_assessment_failure_rolls_back_existing_kol_update(client):
    upload_csv(
        client,
        "platform,country,handle,name,followers\nYouTube,UK,Existing,Existing,123\n",
    )
    workbook = Workbook()
    commercial = workbook.active
    commercial.title = "Commercial Value Model"
    commercial.append([
        "KOLName", "Primary Market (DE/GB/FR/MULTI)", "Content Direction (review/ev/luxury/family/tech/lifestyle)",
    ])
    commercial.append(["Existing", "GB", "ev"])
    risk = workbook.create_sheet("Risk Assessment Model")
    risk.append(["KOLName", "Primary Market (DE/GB/FR/MULTI)", "Fake Follower Share %"])
    risk.append(["Existing", "GB", "not-a-number"])
    content = BytesIO()
    workbook.save(content)

    response = client.post("/imports", files={"file": ("atomic-update.xlsx", content.getvalue())})

    assert response.json()["updated"] == 0
    assert response.json()["failed"] == 1
    existing = client.get("/kols", params={"keyword": "Existing"}).json()[0]
    assert existing["country"] == "GB"
    assert existing["content_categories"] is None
    assert existing["followers"] == 123
def test_import_deduplicates_canonical_youtube_url_and_handle_case(client):
    first = "platform,country,handle,profile_url,followers\nYouTube,GB,@Case,https://www.youtube.com/@Case/,100\n"
    second = "platform,country,handle,profile_url,followers\nyoutube,UK,@other,http://youtube.com/@case?x=1,200\n"
    assert upload_csv(client, first).json()["created"] == 1
    assert upload_csv(client, second).json()["updated"] == 1
    assert len(client.get("/kols").json()) == 1
